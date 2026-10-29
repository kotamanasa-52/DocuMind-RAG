# server.py
import os
import uuid
from pathlib import Path

from flask import Flask, jsonify, render_template, request
from werkzeug.utils import secure_filename

from langchain_core.documents import Document
from langchain_core.prompts import ChatPromptTemplate
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.embeddings import OllamaEmbeddings
from langchain_community.vectorstores import Chroma
from langchain_ollama import ChatOllama

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024  # 20 MB

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)

ALLOWED_EXTENSIONS = {"pdf", "txt", "docx", "csv", "xlsx", "xls"}

# One active document per running server.
active_vectorstore = None
active_filename = None

embeddings = OllamaEmbeddings(model="nomic-embed-text")
llm = ChatOllama(model="llama3.2", temperature=0)


def allowed_file(filename):
    return (
        "." in filename
        and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS
    )


def load_document(file_path, extension):
    """Load supported document types into LangChain Document objects."""
    if extension == "pdf":
        from langchain_community.document_loaders import PyPDFLoader
        return PyPDFLoader(str(file_path)).load()

    if extension == "txt":
        from langchain_community.document_loaders import TextLoader
        return TextLoader(str(file_path), encoding="utf-8").load()

    if extension == "docx":
        from langchain_community.document_loaders import Docx2txtLoader
        return Docx2txtLoader(str(file_path)).load()

    if extension == "csv":
        from langchain_community.document_loaders import CSVLoader
        return CSVLoader(str(file_path), encoding="utf-8").load()

    if extension == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(str(file_path), data_only=True, read_only=True)
        documents = []

        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value is not None]
                if values:
                    rows.append(" | ".join(values))

            if rows:
                documents.append(
                    Document(
                        page_content="\n".join(rows),
                        metadata={"source": file_path.name, "sheet": sheet.title},
                    )
                )

        workbook.close()
        return documents

    if extension == "xls":
        import xlrd

        workbook = xlrd.open_workbook(str(file_path))
        documents = []

        for sheet in workbook.sheets():
            rows = []
            for row_index in range(sheet.nrows):
                values = [
                    str(value)
                    for value in sheet.row_values(row_index)
                    if str(value).strip()
                ]
                if values:
                    rows.append(" | ".join(values))

            if rows:
                documents.append(
                    Document(
                        page_content="\n".join(rows),
                        metadata={"source": file_path.name, "sheet": sheet.name},
                    )
                )

        return documents

    raise ValueError("This file type is not supported.")


@app.route("/")
def home():
    return render_template("index.html")


@app.route("/api/upload", methods=["POST"])
def upload_document():
    global active_vectorstore, active_filename

    uploaded_file = request.files.get("document")

    if not uploaded_file or not uploaded_file.filename:
        return jsonify({"error": "Please choose a document first."}), 400

    safe_name = secure_filename(uploaded_file.filename)

    if not safe_name or not allowed_file(safe_name):
        return jsonify({
            "error": "Unsupported file type. Upload PDF, TXT, DOCX, CSV, XLSX, or XLS."
        }), 400

    extension = safe_name.rsplit(".", 1)[1].lower()
    saved_path = UPLOAD_DIR / f"{uuid.uuid4().hex}_{safe_name}"

    try:
        uploaded_file.save(saved_path)

        documents = load_document(saved_path, extension)
        documents = [
            document for document in documents
            if document.page_content and document.page_content.strip()
        ]

        if not documents:
            return jsonify({
                "error": "No readable text was found in this document."
            }), 400

        splitter = RecursiveCharacterTextSplitter(
            chunk_size=500,
            chunk_overlap=50
        )
        chunks = splitter.split_documents(documents)

        if not chunks:
            return jsonify({
                "error": "This document could not be divided into text sections."
            }), 400

        # Create a fresh vector store for the currently uploaded document.
        active_vectorstore = Chroma.from_documents(
            documents=chunks,
            embedding=embeddings,
            collection_name=f"documind_{uuid.uuid4().hex}",
        )
        active_filename = safe_name

        return jsonify({
            "message": "Document processed successfully.",
            "filename": safe_name,
            "chunks": len(chunks),
        })

    except Exception as exc:
        app.logger.exception("Document processing failed")
        return jsonify({
            "error": (
                "Could not process this document. Check that Ollama is running, "
                "the required models are installed, and the file is readable. "
                f"Details: {str(exc)[:250]}"
            )
        }), 500

    finally:
        # Remove the temporary uploaded file after processing.
        try:
            if saved_path.exists():
                saved_path.unlink()
        except OSError:
            pass


@app.route("/api/ask", methods=["POST"])
def ask_question():
    if active_vectorstore is None:
        return jsonify({"error": "Upload and process a document first."}), 400

    data = request.get_json(silent=True) or {}
    question = str(data.get("question", "")).strip()

    if not question:
        return jsonify({"error": "Enter a question first."}), 400

    try:
        matches = active_vectorstore.similarity_search(question, k=4)

        if not matches:
            return jsonify({
                "answer": "I could not find relevant information in this document.",
                "sources": [],
                "filename": active_filename,
            })

        context_parts = []
        sources = []

        for index, document in enumerate(matches, start=1):
            metadata = document.metadata or {}
            location = metadata.get("page")
            sheet = metadata.get("sheet")

            if location is not None:
                location_text = f"Page {int(location) + 1}"
            elif sheet:
                location_text = f"Sheet: {sheet}"
            else:
                location_text = f"Section {index}"

            context_parts.append(
                f"[Source {index}: {location_text}]\n{document.page_content}"
            )
            sources.append({
                "title": active_filename or "Uploaded document",
                "location": location_text,
                "text": document.page_content[:700],
            })

        context = "\n\n".join(context_parts)

        prompt = ChatPromptTemplate.from_messages([
            (
                "system",
                """You are DocuMind, a helpful document question-answering assistant.

Answer the user's question using only the document context provided.
If the context does not contain the answer, clearly say that the document
does not provide enough information. Do not invent facts.
Use readable paragraphs and bullet points when they improve clarity.
You may refer to the source labels when useful.

DOCUMENT CONTEXT:
{context}"""
            ),
            ("human", "{question}"),
        ])

        response = (prompt | llm).invoke({
            "context": context,
            "question": question,
        })

        return jsonify({
            "answer": response.content,
            "sources": sources,
            "filename": active_filename,
        })

    except Exception as exc:
        app.logger.exception("Question answering failed")
        return jsonify({
            "error": (
                "Could not generate an answer. Check that Ollama is running "
                f"and the required models are installed. Details: {str(exc)[:250]}"
            )
        }), 500


@app.errorhandler(413)
def file_too_large(_error):
    return jsonify({"error": "The file is too large. Maximum size is 20 MB."}), 413


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=True)